#!/usr/bin/env python3
"""
Builds the processed data/*.json and data/*.geojson files consumed by the web
map (public/app.js) from the raw source package in data/source/.

Inputs (data/source/):
  - atlanta_chevy_gis_package.xlsx  (Dealers, ZIP_Data, County_Summary,
    Demographics_Sample tabs -- see PROJECT_BRIEF.md)
  - dealers.geojson                 (point layer, 19 dealers)
  - boundaries/ga_zcta_raw.json     (GA ZCTA polygons, TIGER/Line-derived,
    STATEFP10/ZCTA5CE10 fields -- source: OpenDataDE/State-zip-code-GeoJSON)
  - boundaries/us_counties_raw.json (US county polygons, STATE/COUNTY/NAME
    fields -- source: plotly/datasets geojson-counties-fips.json, itself
    derived from Census cartographic boundary files)

Outputs (public/data/, fetched directly by the web app):
  - dealers.geojson          dealer points + attributes, radius_miles kept
  - zip_choropleth.geojson   ZCTA polygons for the 334 buyer ZIPs, joined to
                             registration + WUPA viewership
  - county_choropleth.geojson  county polygons for the 59 counties in
                             County_Summary, joined to registration + WUPA
  - demographics.json        the one real ACS ZIP profile (30291), keyed by
                             ZIP so more can be appended later without
                             touching app.js
  - zip_data_full.json       all 334 ZIP_Data rows (used for search/lookup;
                             includes the 42 rows with no county match)

Re-run with: python3 scripts/build_data.py
"""
import json
from pathlib import Path

import openpyxl
from shapely.geometry import shape, mapping
from shapely import simplify

# The association's footprint reaches across the Atlanta DMA into edge
# counties of neighboring states, so County_Summary isn't Georgia-only.
STATE_FIPS = {
    "Alabama": "01", "Georgia": "13", "North Carolina": "37",
    "South Carolina": "45", "Tennessee": "47", "Florida": "12",
}

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "source"
OUT = ROOT / "public" / "data"
OUT.mkdir(parents=True, exist_ok=True)

# Simplification tolerance in degrees. ~0.0003 deg is roughly 30m at this
# latitude -- fine for a choropleth at metro-area zoom, and cuts the ZCTA
# file from ~30MB of raw TIGER precision down to a browser-friendly size.
SIMPLIFY_TOLERANCE = 0.0003


def simplify_geom(geom_dict):
    geom = shape(geom_dict)
    simplified = simplify(geom, SIMPLIFY_TOLERANCE, preserve_topology=True)
    return mapping(simplified)


def load_xlsx():
    wb = openpyxl.load_workbook(SRC / "atlanta_chevy_gis_package.xlsx", data_only=True)

    def rows(sheet_name, key_ok):
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        for row in it:
            if row[0] is None or not key_ok(row[0]):
                continue  # skip blank spacer rows, TOTAL rows, and footnotes
            yield dict(zip(header, row))

    return {
        "dealers": list(rows("Dealers", lambda k: True)),
        "zip_data": list(rows("ZIP_Data", lambda k: str(k).isdigit())),
        "county_summary": list(rows("County_Summary", lambda k: ", " in str(k))),
        "demographics_sample": list(wb["Demographics_Sample"].iter_rows(values_only=True)),
    }


def build_dealers_geojson(sheet_rows):
    """dealers.geojson supplies verified point geometry + place_id; the xlsx
    Dealers tab is the source of truth for registration volume (the uploaded
    geojson is missing YTD figures for several dealers that the xlsx has)."""
    by_place_id = {r["Place ID"]: r for r in sheet_rows}

    src_path = SRC / "dealers.geojson"
    data = json.loads(src_path.read_text())
    for feat in data["features"]:
        props = feat["properties"]
        xlsx_row = by_place_id.get(props.get("place_id"))
        if xlsx_row:
            props["veh_ytd_2026"] = xlsx_row["Rolling YTD 2026 (units)"]
            props["veh_ytd_2025"] = xlsx_row["Rolling YTD 2025 (units)"]
            props["yoy_delta"] = xlsx_row["YoY Delta"]
        else:
            v26, v25 = props.get("veh_ytd_2026"), props.get("veh_ytd_2025")
            props["yoy_delta"] = v26 - v25 if isinstance(v26, (int, float)) and isinstance(v25, (int, float)) else None
    (OUT / "dealers.geojson").write_text(json.dumps(data, indent=2))
    return len(data["features"])


def build_zip_choropleth(zip_rows):
    by_zip = {str(r["ZIP Code"]).zfill(5): r for r in zip_rows}

    raw = json.loads((SRC / "boundaries" / "ga_zcta_raw.json").read_text())
    matched = {}
    for feat in raw["features"]:
        zcta = feat["properties"].get("ZCTA5CE10") or feat["properties"].get("ZCTA5CE20")
        if zcta in by_zip and zcta not in matched:
            matched[zcta] = feat

    features = []
    for zcta, row in by_zip.items():
        feat = matched.get(zcta)
        if feat is None:
            continue  # boundary not found in TIGER extract; still present in zip_data_full.json
        features.append({
            "type": "Feature",
            "properties": {
                "zip": zcta,
                "veh_ytd_2026": row["Vehicles YTD 2026"],
                "veh_ytd_2025": row["Vehicles YTD 2025"],
                "yoy_delta": row["YoY Delta"],
                "wupa_avg_imp": row["WUPA Avg News Imp (P2+)"],
                "wupa_tier": row["WUPA Viewing Tier"],
                "leading_dealer": row["Leading Dealer in ZIP"],
            },
            "geometry": simplify_geom(feat["geometry"]),
        })

    fc = {"type": "FeatureCollection", "name": "zip_choropleth", "features": features}
    (OUT / "zip_choropleth.geojson").write_text(json.dumps(fc))
    return len(features), len(by_zip)


def build_county_choropleth(county_rows):
    def split_county_state(full_name):
        county_part, _, state_part = full_name.partition(", ")
        return county_part.replace(" County", "").strip(), state_part.strip()

    by_key = {}
    for r in county_rows:
        name, state = split_county_state(r["County"])
        by_key[(name, state)] = r

    raw = json.loads((SRC / "boundaries" / "us_counties_raw.json").read_text())
    fips_to_state = {v: k for k, v in STATE_FIPS.items()}
    matched = {}
    for feat in raw["features"]:
        props = feat["properties"]
        state = fips_to_state.get(props.get("STATE"))
        if state:
            matched[(props["NAME"], state)] = feat

    features = []
    missing = []
    for (name, state), row in by_key.items():
        feat = matched.get((name, state))
        if feat is None:
            missing.append(f"{name} County, {state}")
            continue
        features.append({
            "type": "Feature",
            "properties": {
                "county": name,
                "state": state,
                "veh_ytd_2026": row["Vehicles YTD 2026 (Polk)"],
                "wupa_total_avg_imp": row["WUPA Total Avg News Imp (P2+)"],
                "zip_count": row["ZIP Codes in County (mapped)"],
            },
            "geometry": simplify_geom(feat["geometry"]),
        })

    fc = {"type": "FeatureCollection", "name": "county_choropleth", "features": features}
    (OUT / "county_choropleth.geojson").write_text(json.dumps(fc))
    return len(features), len(by_key), missing


def build_demographics(demo_rows):
    header_line = demo_rows[0][0]  # "ZIP 30291 -- Union City, GA (Fulton County)"
    zip_code = "".join(ch for ch in header_line.split()[1] if ch.isdigit())
    metrics = {}
    for row in demo_rows[3:]:
        if row[0] is None:
            continue
        metrics[row[0]] = row[1]
    out = {zip_code: {"label": header_line, "metrics": metrics}}
    (OUT / "demographics.json").write_text(json.dumps(out, indent=2))
    return zip_code


def build_zip_data_full(zip_rows):
    out = [
        {
            "zip": str(r["ZIP Code"]).zfill(5),
            "veh_ytd_2026": r["Vehicles YTD 2026"],
            "veh_ytd_2025": r["Vehicles YTD 2025"],
            "yoy_delta": r["YoY Delta"],
            "wupa_avg_imp": r["WUPA Avg News Imp (P2+)"],
            "wupa_tier": r["WUPA Viewing Tier"],
            "leading_dealer": r["Leading Dealer in ZIP"],
        }
        for r in zip_rows
    ]
    (OUT / "zip_data_full.json").write_text(json.dumps(out, indent=2))
    return len(out)


def main():
    sheets = load_xlsx()

    n_dealers = build_dealers_geojson(sheets["dealers"])
    print(f"dealers.geojson: {n_dealers} dealers")

    n_zip_matched, n_zip_total = build_zip_choropleth(sheets["zip_data"])
    print(f"zip_choropleth.geojson: {n_zip_matched}/{n_zip_total} ZIPs matched to boundaries")

    n_county_matched, n_county_total, missing = build_county_choropleth(sheets["county_summary"])
    print(f"county_choropleth.geojson: {n_county_matched}/{n_county_total} counties matched to boundaries")
    if missing:
        print(f"  unmatched counties: {missing}")

    zip_code = build_demographics(sheets["demographics_sample"])
    print(f"demographics.json: sample ZIP {zip_code}")

    n_full = build_zip_data_full(sheets["zip_data"])
    print(f"zip_data_full.json: {n_full} rows")


if __name__ == "__main__":
    main()
